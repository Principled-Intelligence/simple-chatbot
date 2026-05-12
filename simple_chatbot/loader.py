from dataclasses import dataclass, field
from pathlib import Path

from loguru import logger

from simple_chatbot.config import SimpleChatbotConfig

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


@dataclass
class Document:
    text: str
    metadata: dict = field(default_factory=dict)


def _load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = list(reader.pages)
    logger.bind(path=path.name, page_count=len(pages)).debug("PDF loaded")
    return "\n".join(page.extract_text() or "" for page in pages)


def _load_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    paragraphs = [p.text for p in doc.paragraphs]
    logger.bind(path=path.name, paragraph_count=len(paragraphs)).debug("DOCX loaded")
    return "\n".join(paragraphs)


def _chunk(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
    if len(text) <= chunk_size:
        return [text]
    stride = chunk_size - chunk_overlap
    return [text[i : i + chunk_size] for i in range(0, len(text) - chunk_overlap, stride)]


def load_documents(config: SimpleChatbotConfig) -> list[Document]:
    docs_dir = Path(config.docs_dir)
    logger.bind(docs_dir=str(docs_dir.resolve())).info("Scanning documents directory")

    all_files = [p for p in sorted(docs_dir.rglob("*")) if p.is_file()]
    supported = [p for p in all_files if p.suffix.lower() in SUPPORTED_EXTENSIONS]
    skipped = len(all_files) - len(supported)

    logger.bind(
        total=len(all_files),
        supported=len(supported),
        skipped=skipped,
    ).info("Document scan complete")

    docs: list[Document] = []

    for path in supported:
        ext = path.suffix.lower()
        logger.bind(extension=ext, path=str(path)).debug("Loading document file")

        try:
            if ext == ".pdf":
                raw = _load_pdf(path)
            elif ext == ".docx":
                raw = _load_docx(path)
            else:
                raw = _load_txt(path)
        except Exception as e:
            logger.bind(path=str(path), error=str(e)).warning("Could not read document file")
            continue

        raw = raw.strip()
        if not raw:
            logger.bind(path=str(path)).warning("Document file is empty after stripping")
            continue

        chunks = _chunk(raw, config.chunk_size, config.chunk_overlap)
        source = str(path.relative_to(docs_dir))
        logger.bind(
            source=source,
            chars=len(raw),
            chunk_count=len(chunks),
            size=config.chunk_size,
            overlap=config.chunk_overlap,
        ).debug("Document chunked")

        for i, chunk in enumerate(chunks):
            docs.append(Document(text=chunk, metadata={"source": source, "chunk": i}))

    logger.bind(
        files=len(supported),
        chunks=len(docs),
    ).info("Document loading complete")
    return docs
